"""Comprehensive verification test suite for Resume ATS Checker."""

import io
import sys
from pathlib import Path

# Add src to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from pypdf import PdfWriter
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches

from resume_ats_checker.parsers.document_parser import parse_document
from resume_ats_checker.database.connection import check_connection, init_db
from resume_ats_checker.database.repository import (
    save_evaluation,
    get_evaluation_by_id,
    get_recent_evaluations,
    update_suggested_resume,
    update_cover_letter,
)


def test_pdf_parsing():
    print("Testing PDF Parser...")
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    # Write blank or sample pdf
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()
    text, meta = parse_document("sample_resume.pdf", pdf_bytes)
    assert meta["format"] == "pdf", f"Expected format pdf, got {meta['format']}"
    print("  ✓ PDF parser executed successfully.")


def test_docx_parsing():
    print("Testing DOCX Parser...")
    doc = DocxDocument()
    doc.add_heading("John Doe - Senior AI Engineer", level=1)
    doc.add_paragraph("Experienced with Python, LangChain, PostgreSQL, and LLMs.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    text, meta = parse_document("resume.docx", docx_bytes)
    assert "Senior AI Engineer" in text, "Failed to find heading text in docx extraction"
    assert meta["format"] == "docx"
    print("  ✓ DOCX parser extracted text properly.")


def test_pptx_parsing():
    print("Testing PPTX Parser...")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Candidate Portfolio & Skills"
    subtitle.text = "Python, Streamlit, LangChain, OpenAI"
    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    text, meta = parse_document("portfolio.pptx", pptx_bytes)
    assert "Candidate Portfolio & Skills" in text, "Failed to extract title from PPTX"
    assert meta["format"] == "pptx"
    print("  ✓ PPTX parser extracted presentation text properly.")


def test_database_integration():
    print("Testing PostgreSQL Connection and Tables...")
    connected, msg = check_connection()
    print(f"  Connection check: connected={connected}, msg={msg}")
    if connected:
        init_ok, init_msg = init_db()
        print(f"  Database init: success={init_ok}, msg={init_msg}")
        assert init_ok, f"Init DB failed: {init_msg}"

        # Test Save & Retrieve
        test_id = "test-eval-001"
        save_res = save_evaluation({
            "id": test_id,
            "candidate_name": "Test Candidate",
            "resume_filename": "test_resume.pdf",
            "resume_format": "pdf",
            "job_title": "AI Platform Engineer",
            "job_description": "Seeking an AI Engineer with LangChain and PostgreSQL skills.",
            "resume_text": "Experienced engineer with LangChain and SQL.",
            "overall_score": 88,
            "skills_score": 90,
            "experience_score": 85,
            "formatting_score": 92,
            "summary": "Strong candidate match.",
            "missing_keywords": ["Docker", "Kubernetes"],
            "strengths": ["Deep LangChain expertise"],
            "weaknesses": ["Container orchestration not mentioned"],
            "suggestions": ["Add Docker deployment metrics"],
        })
        assert save_res, "Failed to save evaluation to PostgreSQL"

        record = get_evaluation_by_id(test_id)
        assert record is not None, "Failed to retrieve evaluation by ID"
        assert record["overall_score"] == 88, f"Expected score 88, got {record['overall_score']}"

        # Test updates
        update_suggested_resume(test_id, "# Rewritten Resume\n\n- Accomplished X...")
        update_cover_letter(test_id, "Dear Hiring Manager...")
        rec_updated = get_evaluation_by_id(test_id)
        assert "Rewritten Resume" in rec_updated["suggested_resume"]
        assert "Dear Hiring Manager" in rec_updated["cover_letter"]

        recent = get_recent_evaluations(limit=3)
        assert len(recent) > 0, "Expected recent evaluations list to be non-empty"
        print("  ✓ PostgreSQL tables, CRUD operations, and JSON serialization verified!")
    else:
        print("  ⚠️ PostgreSQL not reachable with current credentials (skipping live CRUD).")


from resume_ats_checker.rag.builder_chain import generate_default_sample_resume, StructuredResume


def test_builder_schema():
    print("Testing Structured Resume Schema...")
    sample = generate_default_sample_resume()
    assert isinstance(sample, StructuredResume)
    assert sample.full_name == "Debaranjan"
    assert len(sample.work_experience) >= 1
    assert sample.work_experience[0].company == "Capgemini"
    assert len(sample.work_experience[0].bullets) >= 5
    assert len(sample.education) >= 1
    print("  ✓ Resume Builder schema and default generation verified.")


if __name__ == "__main__":
    test_pdf_parsing()
    test_docx_parsing()
    test_pptx_parsing()
    test_database_integration()
    test_builder_schema()
    print("\n🎉 ALL TESTS PASSED SUCCESSFULLY!")

