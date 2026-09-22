"""Resume document exporter for Word (.docx), PDF (.pdf), and Plain Text (.txt)."""

import io
import re
from typing import Dict, Any, Optional
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.units import inch

from resume_ats_checker.rag.builder_chain import StructuredResume


# ==============================================================================
# 1. Plain Text Exporter
# ==============================================================================

def generate_txt_from_structured_resume(resume: StructuredResume, selections: Dict[str, Any]) -> str:
    """Compile structured resume into clean, ATS-readable plain text."""
    lines = []
    lines.append(resume.full_name.upper())
    contact_parts = []
    if resume.target_title:
        contact_parts.append(resume.target_title)
    if resume.location:
        contact_parts.append(resume.location)
    if resume.email:
        contact_parts.append(resume.email)
    if resume.phone:
        contact_parts.append(resume.phone)
    if resume.linkedin:
        contact_parts.append(resume.linkedin)
    if resume.github_portfolio:
        contact_parts.append(resume.github_portfolio)

    lines.append(" | ".join(contact_parts))
    lines.append("-" * 60)
    lines.append("")

    if selections.get("include_summary", True) and resume.professional_summary:
        lines.append("PROFESSIONAL SUMMARY")
        lines.append(resume.professional_summary)
        lines.append("")

    if resume.skills_languages or resume.skills_frameworks or resume.skills_cloud_tools:
        lines.append("TECHNICAL SKILLS")
        if resume.skills_languages:
            lines.append(f"• Languages: {', '.join(resume.skills_languages)}")
        if resume.skills_frameworks:
            lines.append(f"• Frameworks & Platforms: {', '.join(resume.skills_frameworks)}")
        if resume.skills_cloud_tools:
            lines.append(f"• Cloud & Tools: {', '.join(resume.skills_cloud_tools)}")
        lines.append("")

    has_exp = False
    exp_lines = ["PROFESSIONAL EXPERIENCE"]
    for exp_idx, exp in enumerate(resume.work_experience):
        if selections.get(f"exp_comp_{exp_idx}", True):
            has_exp = True
            exp_lines.append(f"{exp.company} - {exp.role} ({exp.location} | {exp.start_date} - {exp.end_date})")
            for b_idx, bullet in enumerate(exp.bullets):
                if selections.get(f"b_{exp_idx}_{b_idx}", True):
                    exp_lines.append(f"  • {bullet}")
            exp_lines.append("")
    if has_exp:
        lines.extend(exp_lines)

    has_edu = False
    edu_lines = ["EDUCATION"]
    for edu_idx, edu in enumerate(resume.education):
        if selections.get(f"edu_item_{edu_idx}", True):
            has_edu = True
            edu_lines.append(f"• {edu.degree} - {edu.institution}, {edu.location} ({edu.start_date} - {edu.end_date})")
    if has_edu:
        edu_lines.append("")
        lines.extend(edu_lines)

    has_cert = False
    cert_lines = ["CERTIFICATIONS"]
    for c_idx, cert in enumerate(resume.certifications):
        if selections.get(f"cert_{c_idx}", True):
            has_cert = True
            cert_lines.append(f"• {cert}")
    if has_cert:
        cert_lines.append("")
        lines.extend(cert_lines)

    has_proj = False
    proj_lines = ["KEY PROJECTS"]
    for p_idx, proj in enumerate(resume.projects):
        if selections.get(f"proj_{p_idx}", True):
            has_proj = True
            proj_lines.append(f"{proj.title} ({proj.technologies})")
            if proj.description:
                proj_lines.append(f"  {proj.description}")
            for b in proj.bullets:
                proj_lines.append(f"  • {b}")
            proj_lines.append("")
    if has_proj:
        lines.extend(proj_lines)

    return "\n".join(lines)


# ==============================================================================
# 2. Microsoft Word (.docx) Exporter
# ==============================================================================

def generate_docx_from_structured_resume(resume: StructuredResume, selections: Dict[str, Any]) -> bytes:
    """Generate a professionally formatted, ATS-compliant Microsoft Word (.docx) resume."""
    doc = DocxDocument()

    # Configure clean 0.75-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Helper styling function
    def add_section_heading(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(30, 41, 59)  # Dark slate

    # 1. Header: Name & Contact Info
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run(resume.full_name)
    name_run.bold = True
    name_run.font.name = "Calibri"
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = RGBColor(15, 23, 42)

    # Target Title & Contact details
    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(12)

    contact_parts = []
    if resume.target_title:
        contact_parts.append(resume.target_title)
    if resume.location:
        contact_parts.append(resume.location)
    if resume.email:
        contact_parts.append(resume.email)
    if resume.phone:
        contact_parts.append(resume.phone)
    if resume.linkedin:
        contact_parts.append(resume.linkedin)

    c_run = contact_p.add_run("  |  ".join(contact_parts))
    c_run.font.name = "Calibri"
    c_run.font.size = Pt(9.5)
    c_run.font.color.rgb = RGBColor(71, 85, 105)

    # 2. Professional Summary
    if selections.get("include_summary", True) and resume.professional_summary:
        add_section_heading("Professional Summary")
        sum_p = doc.add_paragraph()
        sum_p.paragraph_format.space_after = Pt(6)
        sum_run = sum_p.add_run(resume.professional_summary)
        sum_run.font.name = "Calibri"
        sum_run.font.size = Pt(10.5)

    # 3. Technical Skills
    if resume.skills_languages or resume.skills_frameworks or resume.skills_cloud_tools:
        add_section_heading("Technical Skills")
        if resume.skills_languages:
            sk_p = doc.add_paragraph(style="List Bullet")
            sk_p.paragraph_format.space_after = Pt(2)
            lbl = sk_p.add_run("Languages: ")
            lbl.bold = True
            lbl.font.name = "Calibri"
            lbl.font.size = Pt(10.5)
            val = sk_p.add_run(", ".join(resume.skills_languages))
            val.font.name = "Calibri"
            val.font.size = Pt(10.5)

        if resume.skills_frameworks:
            sk_p = doc.add_paragraph(style="List Bullet")
            sk_p.paragraph_format.space_after = Pt(2)
            lbl = sk_p.add_run("Frameworks & Platforms: ")
            lbl.bold = True
            lbl.font.name = "Calibri"
            lbl.font.size = Pt(10.5)
            val = sk_p.add_run(", ".join(resume.skills_frameworks))
            val.font.name = "Calibri"
            val.font.size = Pt(10.5)

        if resume.skills_cloud_tools:
            sk_p = doc.add_paragraph(style="List Bullet")
            sk_p.paragraph_format.space_after = Pt(4)
            lbl = sk_p.add_run("Cloud & Tools: ")
            lbl.bold = True
            lbl.font.name = "Calibri"
            lbl.font.size = Pt(10.5)
            val = sk_p.add_run(", ".join(resume.skills_cloud_tools))
            val.font.name = "Calibri"
            val.font.size = Pt(10.5)

    # 4. Work Experience
    has_exp = any(selections.get(f"exp_comp_{i}", True) for i in range(len(resume.work_experience)))
    if has_exp:
        add_section_heading("Professional Experience")
        for exp_idx, exp in enumerate(resume.work_experience):
            if selections.get(f"exp_comp_{exp_idx}", True):
                # Job Header line
                jh_p = doc.add_paragraph()
                jh_p.paragraph_format.space_before = Pt(4)
                jh_p.paragraph_format.space_after = Pt(2)

                comp_run = jh_p.add_run(exp.company)
                comp_run.bold = True
                comp_run.font.name = "Calibri"
                comp_run.font.size = Pt(11)

                role_run = jh_p.add_run(f" — {exp.role}")
                role_run.font.name = "Calibri"
                role_run.font.size = Pt(11)

                dates_run = jh_p.add_run(f"  ({exp.location} | {exp.start_date} – {exp.end_date})")
                dates_run.italic = True
                dates_run.font.name = "Calibri"
                dates_run.font.size = Pt(10)
                dates_run.font.color.rgb = RGBColor(100, 116, 139)

                # Bullets
                for b_idx, bullet in enumerate(exp.bullets):
                    if selections.get(f"b_{exp_idx}_{b_idx}", True):
                        bp = doc.add_paragraph(style="List Bullet")
                        bp.paragraph_format.space_after = Pt(2)
                        brun = bp.add_run(bullet)
                        brun.font.name = "Calibri"
                        brun.font.size = Pt(10)

    # 5. Education
    has_edu = any(selections.get(f"edu_item_{i}", True) for i in range(len(resume.education)))
    if has_edu:
        add_section_heading("Education")
        for edu_idx, edu in enumerate(resume.education):
            if selections.get(f"edu_item_{edu_idx}", True):
                ed_p = doc.add_paragraph(style="List Bullet")
                ed_p.paragraph_format.space_after = Pt(2)
                deg_run = ed_p.add_run(f"{edu.degree}")
                deg_run.bold = True
                deg_run.font.name = "Calibri"
                deg_run.font.size = Pt(10.5)

                inst_run = ed_p.add_run(f" — {edu.institution}, {edu.location} ({edu.start_date} – {edu.end_date})")
                inst_run.font.name = "Calibri"
                inst_run.font.size = Pt(10)

    # 6. Certifications
    has_cert = any(selections.get(f"cert_{i}", True) for i in range(len(resume.certifications)))
    if has_cert:
        add_section_heading("Certifications")
        for c_idx, cert in enumerate(resume.certifications):
            if selections.get(f"cert_{c_idx}", True):
                cp = doc.add_paragraph(style="List Bullet")
                cp.paragraph_format.space_after = Pt(2)
                c_run = cp.add_run(cert)
                c_run.font.name = "Calibri"
                c_run.font.size = Pt(10)

    # 7. Projects
    has_proj = any(selections.get(f"proj_{i}", True) for i in range(len(resume.projects)))
    if has_proj:
        add_section_heading("Key Projects")
        for p_idx, proj in enumerate(resume.projects):
            if selections.get(f"proj_{p_idx}", True):
                pp = doc.add_paragraph()
                pp.paragraph_format.space_after = Pt(1)
                p_run = pp.add_run(proj.title)
                p_run.bold = True
                p_run.font.name = "Calibri"
                p_run.font.size = Pt(10.5)

                if proj.technologies:
                    t_run = pp.add_run(f" ({proj.technologies})")
                    t_run.italic = True
                    t_run.font.name = "Calibri"
                    t_run.font.size = Pt(10)

                for b in proj.bullets:
                    pbp = doc.add_paragraph(style="List Bullet")
                    pbp.paragraph_format.space_after = Pt(2)
                    prun = pbp.add_run(b)
                    prun.font.name = "Calibri"
                    prun.font.size = Pt(10)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ==============================================================================
# 3. PDF (.pdf) Exporter using ReportLab
# ==============================================================================

def generate_pdf_from_structured_resume(resume: StructuredResume, selections: Dict[str, Any]) -> bytes:
    """Generate a clean, high-precision ATS PDF resume using ReportLab Platypus."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )

    styles = getSampleStyleSheet()

    # Custom typography styles
    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=1,  # Center
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=3,
    )

    contact_style = ParagraphStyle(
        "ResumeContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        alignment=1,  # Center
        textColor=colors.HexColor("#475569"),
        spaceAfter=12,
    )

    heading_style = ParagraphStyle(
        "ResumeHeading",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=8,
        spaceAfter=2,
    )

    body_style = ParagraphStyle(
        "ResumeBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=4,
    )

    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.2,
        leading=13,
        leftIndent=14,
        firstLineIndent=-10,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=2.5,
    )

    job_header_style = ParagraphStyle(
        "JobHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=3,
        spaceAfter=2,
    )

    story = []

    # 1. Name & Contact
    story.append(Paragraph(resume.full_name, title_style))
    contact_parts = []
    if resume.target_title:
        contact_parts.append(f"<b>{resume.target_title}</b>")
    if resume.location:
        contact_parts.append(resume.location)
    if resume.email:
        contact_parts.append(resume.email)
    if resume.phone:
        contact_parts.append(resume.phone)
    if resume.linkedin:
        contact_parts.append(resume.linkedin)

    story.append(Paragraph(" &nbsp;|&nbsp; ".join(contact_parts), contact_style))

    def add_pdf_section(title: str):
        story.append(Paragraph(title.upper(), heading_style))
        story.append(HRFlowable(width="100%", thickness=0.75, color=colors.HexColor("#cbd5e1"), spaceBefore=1, spaceAfter=4))

    # 2. Summary
    if selections.get("include_summary", True) and resume.professional_summary:
        add_pdf_section("Professional Summary")
        story.append(Paragraph(resume.professional_summary, body_style))

    # 3. Technical Skills
    if resume.skills_languages or resume.skills_frameworks or resume.skills_cloud_tools:
        add_pdf_section("Technical Skills")
        if resume.skills_languages:
            story.append(Paragraph(f"• <b>Languages:</b> {', '.join(resume.skills_languages)}", bullet_style))
        if resume.skills_frameworks:
            story.append(Paragraph(f"• <b>Frameworks & Platforms:</b> {', '.join(resume.skills_frameworks)}", bullet_style))
        if resume.skills_cloud_tools:
            story.append(Paragraph(f"• <b>Cloud & Tools:</b> {', '.join(resume.skills_cloud_tools)}", bullet_style))

    # 4. Work Experience
    has_exp = any(selections.get(f"exp_comp_{i}", True) for i in range(len(resume.work_experience)))
    if has_exp:
        add_pdf_section("Professional Experience")
        for exp_idx, exp in enumerate(resume.work_experience):
            if selections.get(f"exp_comp_{exp_idx}", True):
                header_text = f"<b>{exp.company}</b> — <i>{exp.role}</i> <font color='#64748b'>({exp.location} | {exp.start_date} – {exp.end_date})</font>"
                story.append(Paragraph(header_text, job_header_style))

                for b_idx, bullet in enumerate(exp.bullets):
                    if selections.get(f"b_{exp_idx}_{b_idx}", True):
                        story.append(Paragraph(f"• {bullet}", bullet_style))
                story.append(Spacer(1, 3))

    # 5. Education
    has_edu = any(selections.get(f"edu_item_{i}", True) for i in range(len(resume.education)))
    if has_edu:
        add_pdf_section("Education")
        for edu_idx, edu in enumerate(resume.education):
            if selections.get(f"edu_item_{edu_idx}", True):
                story.append(Paragraph(f"• <b>{edu.degree}</b> — {edu.institution}, {edu.location} ({edu.start_date} – {edu.end_date})", bullet_style))

    # 6. Certifications
    has_cert = any(selections.get(f"cert_{i}", True) for i in range(len(resume.certifications)))
    if has_cert:
        add_pdf_section("Certifications")
        for c_idx, cert in enumerate(resume.certifications):
            if selections.get(f"cert_{c_idx}", True):
                story.append(Paragraph(f"• {cert}", bullet_style))

    # 7. Projects
    has_proj = any(selections.get(f"proj_{i}", True) for i in range(len(resume.projects)))
    if has_proj:
        add_pdf_section("Key Projects")
        for p_idx, proj in enumerate(resume.projects):
            if selections.get(f"proj_{p_idx}", True):
                story.append(Paragraph(f"<b>{proj.title}</b> ({proj.technologies})", job_header_style))
                if proj.description:
                    story.append(Paragraph(f"<i>{proj.description}</i>", body_style))
                for b in proj.bullets:
                    story.append(Paragraph(f"• {b}", bullet_style))

    doc.build(story)
    return buffer.getvalue()


# ==============================================================================
# 4. Markdown Exporters (For Suggested Complete Resume in ATS Checker)
# ==============================================================================

def generate_docx_from_markdown(markdown_text: str, candidate_name: str = "Candidate") -> bytes:
    """Convert raw Markdown resume text into a formatted Word (.docx) document."""
    doc = DocxDocument()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    for line in markdown_text.splitlines():
        line_s = line.strip()
        if not line_s:
            continue

        if line_s.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line_s[2:].strip())
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif line_s.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line_s[3:].strip().upper())
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif line_s.startswith("- ") or line_s.startswith("• "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            # Simple bold parser for markdown **bold**
            clean_item = line_s[2:].strip()
            parts = re.split(r"(\*\*.*?\*\*)", clean_item)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            parts = re.split(r"(\*\*.*?\*\*)", line_s)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_pdf_from_markdown(markdown_text: str, candidate_name: str = "Candidate") -> bytes:
    """Convert raw Markdown resume text into a high-precision PDF document using ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "MdTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=1,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=4,
    )

    heading_style = ParagraphStyle(
        "MdHeading",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=8,
        spaceAfter=2,
    )

    body_style = ParagraphStyle(
        "MdBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=3,
    )

    bullet_style = ParagraphStyle(
        "MdBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.2,
        leading=13,
        leftIndent=14,
        firstLineIndent=-10,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=2,
    )

    story = []

    for line in markdown_text.splitlines():
        line_s = line.strip()
        if not line_s:
            continue

        # Convert markdown **bold** to <b>bold</b> and *italic* to <i>italic</i> for ReportLab
        formatted_line = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", line_s)
        formatted_line = re.sub(r"\*(.*?)\*", r"<i>\1</i>", formatted_line)

        if line_s.startswith("# "):
            story.append(Paragraph(formatted_line[2:].strip(), title_style))
        elif line_s.startswith("## "):
            story.append(Paragraph(formatted_line[3:].strip().upper(), heading_style))
            story.append(HRFlowable(width="100%", thickness=0.75, color=colors.HexColor("#cbd5e1"), spaceBefore=1, spaceAfter=4))
        elif line_s.startswith("- ") or line_s.startswith("• "):
            story.append(Paragraph(f"• {formatted_line[2:].strip()}", bullet_style))
        else:
            story.append(Paragraph(formatted_line, body_style))

    doc.build(story)
    return buffer.getvalue()
